import os
import io
import csv
import json
import re
import random
import string
from datetime import datetime, date
from typing import List, Optional, Any, Dict
from fastapi import APIRouter, Depends, HTTPException, status, Query, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database import get_db
from app.core.deps import get_current_user, check_permission
from app.models.order import Order
from app.models.inventory import Inventory
from app.models.account import Account
from app.models.user import User
from app.schemas.order import OrderCreate, OrderUpdate, OrderResponse
from app.core.admin_cost_sync import sync_admin_cost_share_for_month
from app.core.url_extractor import extract_product_info_from_url

router = APIRouter(prefix="/orders", tags=["Orders"])

def generate_order_number(db: Session = None):
    while True:
        date_str = datetime.utcnow().strftime("%Y%m%d")
        random_str = ''.join(random.choices(string.ascii_uppercase + string.digits, k=6))
        ord_num = f"ORD-{date_str}-{random_str}"
        if db is None:
            return ord_num
        if not db.query(Order).filter(Order.order_number == ord_num).first():
            return ord_num

from sqlalchemy import or_, func

def populate_order_costs(orders, db: Session):
    is_list = isinstance(orders, list)
    orders_list = orders if is_list else [orders]
    
    for o in orders_list:
        p_cost = getattr(o, 'purchase_cost_inr', 0.0) or 0.0
        a_cost = getattr(o, 'admin_cost_share', 0.0) or 0.0
        o.total_order_cost_inr = p_cost + a_cost
        
    return orders

def get_user_allowed_companies(user: User) -> Optional[List[str]]:
    if not user or user.is_admin:
        return None
    if getattr(user, 'allowed_companies', None):
        return [c.strip() for c in user.allowed_companies.split(',') if c.strip()]
    if user.email == 'ops2@crm.com':
        return ['ADBH', 'Globle', 'Global']
    return None

@router.get("", response_model=List[OrderResponse])
def list_orders(
    status_filter: Optional[str] = None,
    search: Optional[str] = None,
    skip: Optional[int] = Query(None, ge=0, description="Number of items to skip"),
    limit: Optional[int] = Query(None, ge=1, le=1000, description="Max items to return"),
    page: Optional[int] = Query(None, ge=1, description="Page number (1-indexed)"),
    page_size: Optional[int] = Query(None, ge=1, le=1000, description="Items per page"),
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("orders:read"))
):
    query = db.query(Order)
    
    # Filter by user allowed companies (e.g. User 2 ADBH & Global only)
    allowed_comps = get_user_allowed_companies(current_user)
    if allowed_comps:
        query = query.filter(or_(*[Order.company.ilike(f"%{c}%") for c in allowed_comps]))
    else:
        # Filter by specific company/account only if user is an external partner or assigned to a specific account
        is_partner_user = current_user.is_partner or (current_user.role and current_user.role.name == "Channel Partner")
        if not current_user.is_admin and (is_partner_user or current_user.account_name or current_user.account_id):
            comp_name = current_user.account_name
            if comp_name:
                query = query.filter(
                    (Order.company.ilike(f"%{comp_name}%")) |
                    (Order.account_name.ilike(f"%{comp_name}%")) |
                    (Order.seller_account.ilike(f"%{comp_name}%"))
                )
            elif current_user.account_id:
                query = query.filter(Order.account_id == current_user.account_id)

    if status_filter:
        query = query.filter(Order.status == status_filter)
    if search:
        query = query.filter(
            (Order.order_number.ilike(f"%{search}%")) |
            (Order.shipment_id.ilike(f"%{search}%")) |
            (Order.seller_account.ilike(f"%{search}%")) |
            (Order.company.ilike(f"%{search}%")) |
            (Order.consignee_name.ilike(f"%{search}%")) |
            (Order.product_name.ilike(f"%{search}%"))
        )
    
    query = query.order_by(Order.created_at.desc())

    if page is not None and page_size is not None:
        skip = (page - 1) * page_size
        limit = page_size
    if skip is not None:
        query = query.offset(skip)
    if limit is not None:
        query = query.limit(limit)

    orders = query.all()
    populate_order_costs(orders, db)
    return orders

import re

def get_next_shipment_id(db: Session, offset: int = 0) -> str:
    orders = db.query(Order.shipment_id).filter(Order.shipment_id.isnot(None)).all()
    max_num = 0
    prefix = "INBTL"
    digits_len = 3

    for (s_id,) in orders:
        if s_id:
            match = re.search(r'^([A-Za-z\-_]+)(\d+)$', s_id.strip())
            if match:
                prefix = match.group(1)
                digits_len = max(digits_len, len(match.group(2)))
                num = int(match.group(2))
                if num > max_num:
                    max_num = num

    next_num = max_num + 1 + offset
    return f"{prefix}{str(next_num).zfill(digits_len)}"

def get_or_create_inventory_item(
    db: Session,
    product_name: str,
    price_usd: float = 0.0,
    seller_account: Optional[str] = None,
    image_url: Optional[str] = None,
    product_url: Optional[str] = None
) -> Inventory:
    clean_name = product_name.strip() if product_name else ""
    if not clean_name:
        return None

    norm_name = re.sub(r'\s+', ' ', clean_name).strip()

    # Strict case-insensitive, normalized & trimmed search to prevent ANY duplicate product entries in database
    item = db.query(Inventory).filter(
        or_(
            func.lower(func.trim(Inventory.product_name)) == clean_name.lower(),
            func.lower(func.trim(Inventory.product_name)) == norm_name.lower(),
            Inventory.product_name.ilike(clean_name),
            Inventory.product_name.ilike(norm_name)
        )
    ).first()

    if item:
        # Product already exists in database: do NOT create duplicate
        changed = False
        if image_url and (not item.image_url or "unsplash" in item.image_url):
            item.image_url = image_url
            changed = True
        if product_url and product_url.strip():
            clean_url = product_url.strip()
            if not item.product_url or item.product_url.strip() != clean_url:
                item.product_url = clean_url
                changed = True
        if price_usd and (not item.price or item.price == 0):
            item.price = price_usd
            changed = True
        if changed:
            db.commit()
            db.refresh(item)
        return item

    # Product does NOT exist -> Create new unique Inventory item
    prefix_code = "".join(e for e in clean_name if e.isalnum())[:4].upper() or "ITEM"
    rand_suffix = random.randint(100, 999)
    sku_candidate = f"SKU-{prefix_code}-{rand_suffix}"
    while db.query(Inventory).filter(Inventory.sku == sku_candidate).first():
        rand_suffix = random.randint(1000, 9999)
        sku_candidate = f"SKU-{prefix_code}-{rand_suffix}"

    item = Inventory(
        product_name=clean_name,
        price=price_usd or 0.0,
        stock_quantity=0,
        sku=sku_candidate,
        category="General",
        partner_name=seller_account or "General",
        image_url=image_url or "https://images.unsplash.com/photo-1544816155-12df9643f363?w=300",
        product_url=product_url.strip() if product_url else None
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return item

from sqlalchemy.exc import IntegrityError

@router.post("", response_model=OrderResponse)
def create_order(
    order_in: OrderCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("orders:write"))
):
    inventory_item = None
    if order_in.product_id:
        inventory_item = db.query(Inventory).filter(Inventory.id == order_in.product_id).first()

    if not inventory_item and order_in.product_name:
        inventory_item = get_or_create_inventory_item(
            db=db,
            product_name=order_in.product_name,
            price_usd=order_in.price_usd or 0.0,
            seller_account=order_in.seller_account or order_in.company,
            image_url=order_in.product_image,
            product_url=order_in.product_url
        )
    elif inventory_item:
        # Sync product_url and image_url to inventory_item if user provided them
        changed = False
        if order_in.product_url and order_in.product_url.strip():
            clean_url = order_in.product_url.strip()
            if not inventory_item.product_url or inventory_item.product_url.strip() != clean_url:
                inventory_item.product_url = clean_url
                changed = True
        if order_in.product_image and not inventory_item.image_url:
            inventory_item.image_url = order_in.product_image
            changed = True
        if changed:
            db.commit()
            db.refresh(inventory_item)

    # Determine effective product_url: fallback to inventory_item URL if blank
    raw_url = (order_in.product_url or "").strip()
    effective_product_url = raw_url if raw_url else (inventory_item.product_url if inventory_item else None)

    effective_product_image = order_in.product_image or (inventory_item.image_url if inventory_item else None)
    if not effective_product_image and effective_product_url:
        extracted = extract_product_info_from_url(effective_product_url)
        if extracted.get("image_url"):
            effective_product_image = extracted["image_url"]
            if inventory_item and not inventory_item.image_url:
                inventory_item.image_url = effective_product_image
                db.commit()

    account_id = order_in.account_id
    account_name = order_in.account_name

    allowed_comps = get_user_allowed_companies(current_user)
    if allowed_comps:
        comp_val = (order_in.company or "").strip().lower()
        if not any(c.lower() in comp_val for c in allowed_comps):
            raise HTTPException(
                status_code=403,
                detail=f"Not authorized to create orders for company '{order_in.company}'. Allowed: {', '.join(allowed_comps)}"
            )

    is_partner_user = current_user.is_partner or (current_user.role and current_user.role.name == "Channel Partner")
    if not current_user.is_admin and (is_partner_user or current_user.account_name or current_user.account_id):
        account_id = current_user.account_id
        account_name = current_user.account_name

    if not account_id and order_in.seller_account:
        acc = db.query(Account).filter(Account.account_name.ilike(order_in.seller_account.strip())).first()
        if acc:
            account_id = acc.id
            account_name = acc.account_name

    if account_id and not account_name:
        acc = db.query(Account).filter(Account.id == account_id).first()
        if acc:
            account_name = acc.account_name

    if order_in.order_number:
        order_num = order_in.order_number.strip()
        existing_order = db.query(Order).filter(Order.order_number == order_num).first()
        if existing_order:
            raise HTTPException(
                status_code=400,
                detail=f"Order ID '{order_num}' already exists. Please enter a unique Order ID."
            )
    else:
        order_num = generate_order_number(db)

    shipment_num = order_in.shipment_id or get_next_shipment_id(db)

    order = Order(
        order_number=order_num,
        order_process_date=order_in.order_process_date or date.today(),
        last_delivery_date=order_in.last_delivery_date,
        shipping_date=order_in.shipping_date,
        company=order_in.company,
        shipment_id=shipment_num,
        seller_account=order_in.seller_account,
        product_id=inventory_item.id if inventory_item else None,
        product_name=order_in.product_name,
        product_url=effective_product_url,
        product_image=effective_product_image,
        product_items=order_in.product_items,
        qty=order_in.qty,
        product_price=getattr(order_in, 'product_price', None) or order_in.price_usd or 0.0,
        order_status=order_in.order_status or "ADBH",
        purchase_cost_inr=order_in.purchase_cost_inr or 0.0,
        arriving_date=order_in.arriving_date,
        consignee_name=order_in.consignee_name or "Consignee",
        shipment_address_1=order_in.shipment_address_1 or "",
        shipment_address_2=order_in.shipment_address_2 or "",
        city=order_in.city or "",
        state=order_in.state or "",
        zip_code=order_in.zip_code or "",
        mobile_number=order_in.mobile_number or "",
        country=order_in.country or "USA",
        account_id=account_id,
        account_name=account_name,
        status=order_in.status or "Pending",
        delivery_service=order_in.delivery_service,
        shipment_cost=order_in.shipment_cost or 0.0
    )
    try:
        db.add(order)
        db.commit()
        db.refresh(order)
    except IntegrityError as ex:
        db.rollback()
        err_msg = str(ex.orig) if hasattr(ex, 'orig') else str(ex)
        if "order_number" in err_msg or "UNIQUE" in err_msg.upper():
            detail_msg = f"Order with Order ID '{order_num}' already exists."
        else:
            detail_msg = f"Database integrity error: {err_msg}"
        raise HTTPException(
            status_code=400,
            detail=detail_msg
        )
    
    d = order.order_date or order.order_process_date or date.today()
    if d:
        sync_admin_cost_share_for_month(d.strftime("%Y-%m"), db)
        db.refresh(order)
        
    populate_order_costs(order, db)
    return order

import csv
import io
import re
import uuid as uuid_module
from PIL import Image
from fastapi import UploadFile, File
from app.core.s3 import is_s3_enabled, upload_file_to_s3
from app.models.shipment import Shipment


def clean_tracking_id(raw_str: str) -> Optional[str]:
    """Sanitizes and validates tracking numbers from barcodes or text."""
    if not raw_str:
        return None
    raw = str(raw_str).strip()
    
    # Check GS1 pattern with parenthesized Application Identifiers:
    # e.g. (420)61244(92)34690363072206692393 or (420) 61244 (94) 0010...
    m_gs1_parens = re.search(r'\(420\)\s*\d{5}\s*\((9[1-5]|\d{2})\)\s*(\d{18,24})', raw)
    if m_gs1_parens:
        return m_gs1_parens.group(1) + m_gs1_parens.group(2)
        
    # Check general parens GS1 e.g. (92)34690363072206692393
    m_usps_parens = re.search(r'\((9[1-5]|\d{2})\)\s*(\d{18,24})', raw)
    if m_usps_parens:
        return m_usps_parens.group(1) + m_usps_parens.group(2)

    # Remove symbology identifiers like ]C1, ]d2, ]e0 and non-alphanumeric chars
    cleaned = re.sub(r'^\][a-zA-Z0-9]{2}', '', raw)
    cleaned = re.sub(r'[\s\-_()\[\]{}:;,\x1d\x1e\x04]+', '', cleaned)
    
    # 1. GS1-128 barcode format: 420 + 5-digit zip + 20-24 digit USPS tracking
    # e.g. 420612449234690363072206692393
    m_gs1 = re.match(r'^.*?420\d{5}(9[1-5]\d{18,22}|\d{20,24})$', cleaned)
    if m_gs1:
        return m_gs1.group(1)
    
    # 2. USPS 20-24 digit tracking (often starts with 91, 92, 93, 94, 95, etc.)
    m_usps = re.match(r'^(9[1-5]\d{18,22}|\d{20,24})$', cleaned)
    if m_usps:
        return m_usps.group(1)
    
    # 3. UPS tracking: 1Z + 16 alphanumeric characters
    m_ups = re.match(r'^(1Z[A-Z0-9]{16})$', cleaned, re.I)
    if m_ups:
        return m_ups.group(1).upper()
        
    # 4. FedEx tracking: 12, 14, 15, 20, 22 digits
    m_fedex = re.match(r'^(\d{12}|\d{14}|\d{15}|\d{20}|\d{22})$', cleaned)
    if m_fedex:
        return m_fedex.group(1)
        
    # 5. DHL: 10-11 digits or JD...
    m_dhl = re.match(r'^(JD\d{18}|\d{10,11})$', cleaned, re.I)
    if m_dhl:
        return m_dhl.group(1)
        
    # 6. Amazon TBA
    m_amz = re.match(r'^(TBA\d{12})$', cleaned, re.I)
    if m_amz:
        return m_amz.group(1).upper()

    # 7. Generic alphanumeric tracking (e.g. 10 to 34 chars)
    if len(cleaned) >= 10 and re.match(r'^[A-Z0-9]{10,34}$', cleaned, re.I):
        return cleaned

    return None


def extract_tracking_from_text(text: str) -> Optional[str]:
    """Extracts tracking number from raw OCR/PDF text using carrier patterns."""
    if not text:
        return None
        
    # 1. USPS Tracking Header pattern (e.g. USPS TRACKING # USPS Ship \n 9234 6903 6307 2206 6923 93)
    m = re.search(r'USPS\s+TRACKING\s*#?[^\n\r]*[\r\n]+\s*([0-9\s]{20,35})', text, re.I)
    if m:
        val = clean_tracking_id(m.group(1))
        if val:
            return val
            
    # 2. UPS Tracking pattern
    m = re.search(r'\b(1Z\s*[A-Z0-9\s]{16,22})\b', text, re.I)
    if m:
        val = clean_tracking_id(m.group(1))
        if val:
            return val

    # 3. GS1 format with 420 zip prefix in text
    m = re.search(r'\b(420\d{5}[\s\-]*(?:9[1-5][\d\s\-]{18,28}|\d[\d\s\-]{19,30}))\b', text)
    if m:
        val = clean_tracking_id(m.group(1))
        if val:
            return val

    # 4. Spaced 20-24 digit numbers (e.g. 9234 6903 6307 2206 6923 93)
    matches = re.findall(r'\b((?:9[1-5]\d{2}|\d{4})[\s\-]+(?:\d{4}[\s\-]+){3,4}\d{2,6})\b', text)
    for match in matches:
        val = clean_tracking_id(match)
        if val:
            return val

    # 5. Continuous 20-24 digit tracking (USPS / FedEx)
    matches = re.findall(r'\b(9[1-5]\d{18,22}|\d{20,24})\b', text)
    for match in matches:
        val = clean_tracking_id(match)
        if val:
            return val

    # 6. FedEx 12-digit grouped (e.g. 7834 1234 5678)
    matches = re.findall(r'\b(\d{4}[\s\-]+\d{4}[\s\-]+\d{4})\b', text)
    for match in matches:
        val = clean_tracking_id(match)
        if val:
            return val

    # 7. Amazon TBA
    matches = re.findall(r'\b(TBA\d{12})\b', text, re.I)
    if matches:
        return matches[0].upper()

    return None


def extract_tracking_from_barcode_image(img_pil: Image.Image) -> Optional[str]:
    """Scans an image with zxing-cpp to read 1D and 2D barcodes, supporting rotations."""
    try:
        import zxingcpp
        
        # Ensure image is in standard RGB or grayscale mode for zxing-cpp
        img = img_pil
        if img.mode not in ('RGB', 'L', 'RGBA'):
            img = img.convert('RGB')

        results = zxingcpp.read_barcodes(img)
        for r in results:
            if r.text:
                cleaned = clean_tracking_id(r.text)
                if cleaned:
                    return cleaned

        # Try 90, 180, 270 rotations if not detected in default orientation
        for angle in [90, 180, 270]:
            rot_img = img.rotate(angle, expand=True)
            results = zxingcpp.read_barcodes(rot_img)
            for r in results:
                if r.text:
                    cleaned = clean_tracking_id(r.text)
                    if cleaned:
                        return cleaned
    except Exception:
        pass
    return None


def extract_tracking_id(content: bytes, filename: str = "") -> Optional[str]:
    """
    Comprehensive extraction of tracking/forwarding number from PDF or image content.
    Combines text parsing and barcode scanning.
    """
    fn = (filename or "").lower()
    is_pdf = fn.endswith(".pdf") or content[:4] == b"%PDF"

    if is_pdf:
        # 1. Try PDF text layer
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            full_text = ""
            for page in reader.pages:
                try:
                    full_text += (page.extract_text() or "") + "\n"
                except Exception:
                    pass
            
            trk = extract_tracking_from_text(full_text)
            if trk:
                return trk
                
            # 2. Try embedded barcode images in PDF pages
            for page in reader.pages:
                if hasattr(page, 'images'):
                    for img_file in page.images:
                        try:
                            pil_img = Image.open(io.BytesIO(img_file.data))
                            trk = extract_tracking_from_barcode_image(pil_img)
                            if trk:
                                return trk
                        except Exception:
                            pass
        except Exception:
            pass
    else:
        # Try image barcode scan
        try:
            pil_img = Image.open(io.BytesIO(content))
            trk = extract_tracking_from_barcode_image(pil_img)
            if trk:
                return trk
        except Exception:
            pass

    return None


def convert_image_to_pdf_bytes(img_bytes: bytes) -> bytes:
    """Converts image bytes into a single-page PDF with matching dimensions."""
    try:
        from reportlab.pdfgen import canvas
        pil_img = Image.open(io.BytesIO(img_bytes))
        w, h = pil_img.size
        
        pdf_buf = io.BytesIO()
        c = canvas.Canvas(pdf_buf, pagesize=(w, h))
        c.drawInlineImage(pil_img, 0, 0, width=w, height=h)
        c.save()
        return pdf_buf.getvalue()
    except Exception:
        return img_bytes


@router.post("/{order_id}/upload-label")
async def upload_order_label(
    order_id: int,
    file: UploadFile = File(...),
    label_cost_usd: float = 0.0,
    label_free: bool = False,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("orders:write"))
):
    order = db.query(Order).filter(Order.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    content = await file.read()
    orig_name = file.filename or "label.pdf"
    fn_lower = orig_name.lower()
    is_image = fn_lower.endswith(('.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tiff'))

    # Extract tracking / forwarding number from PDF or Image
    extracted_tracking_id = extract_tracking_id(content, orig_name)

    # Convert image to PDF if uploaded as image so viewing & stamping is always a consistent PDF
    upload_content = content
    upload_filename = orig_name
    upload_content_type = file.content_type or "application/pdf"

    if is_image:
        try:
            pdf_bytes = convert_image_to_pdf_bytes(content)
            if pdf_bytes and pdf_bytes[:4] == b"%PDF":
                upload_content = pdf_bytes
                upload_filename = os.path.splitext(orig_name)[0] + ".pdf"
                upload_content_type = "application/pdf"
        except Exception:
            pass

    # Upload the file (S3 or local)
    if is_s3_enabled():
        file_url, _, _ = upload_file_to_s3(
            file_content=upload_content,
            original_filename=upload_filename,
            content_type=upload_content_type,
        )
    else:
        BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
        os.makedirs(UPLOADS_DIR, exist_ok=True)
        unique_name = f"{uuid_module.uuid4().hex[:12]}_{upload_filename.replace(' ', '_')}"
        file_path = os.path.join(UPLOADS_DIR, unique_name)
        with open(file_path, "wb") as buf:
            buf.write(upload_content)
        file_url = f"/uploads/{unique_name}"

    # Update the order
    order.label_pdf_url = file_url
    order.label_cost_usd = label_cost_usd
    order.label_free = label_free
    if extracted_tracking_id:
        order.label_tracking_id = extracted_tracking_id
        
        # Auto-update forwarding_number in any existing shipment for this order
        shipments = db.query(Shipment).filter(Shipment.order_id == order.id).all()
        for s in shipments:
            s.forwarding_number = extracted_tracking_id

    db.commit()
    db.refresh(order)
    populate_order_costs(order, db)

    return {
        "success": True,
        "label_pdf_url": file_url,
        "tracking_id_extracted": extracted_tracking_id,
        "order": order
    }


def stamp_product_name_on_pdf(pdf_bytes: bytes, product_name: str, order_num: str = "", qty: int = 1) -> bytes:
    """
    Stamps product name and order info at the bottom of the PDF label.
    Maintains exact original page dimensions (e.g. 4x6 inches) and scales original content
    proportionately so no part of the shipping label (top or bottom) is ever clipped or cut.
    """
    try:
        from pypdf import PdfReader, PdfWriter, Transformation
        from reportlab.pdfgen import canvas
        from reportlab.lib.colors import HexColor

        # Fallback: if not valid PDF bytes (e.g. legacy image upload), convert first
        if pdf_bytes[:4] != b"%PDF":
            pdf_bytes = convert_image_to_pdf_bytes(pdf_bytes)

        reader = PdfReader(io.BytesIO(pdf_bytes))
        writer = PdfWriter()

        footer_height = 20  # pt for product name footer

        for page in reader.pages:
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)

            # Scale factor so entire label fits perfectly above footer without altering page size
            scale = max(0.85, (height - footer_height) / height)

            # Create background canvas with EXACT same page dimensions
            bg_buf = io.BytesIO()
            bg_can = canvas.Canvas(bg_buf, pagesize=(width, height))

            # Draw clean bottom footer
            bg_can.setFillColor(HexColor("#FFFFFF"))
            bg_can.rect(0, 0, width, footer_height, fill=1, stroke=0)

            # Divider line at the top of footer
            bg_can.setStrokeColor(HexColor("#CBD5E1"))
            bg_can.setLineWidth(0.75)
            bg_can.line(0, footer_height, width, footer_height)

            # Product details text
            bg_can.setFont("Helvetica-Bold", 8)
            bg_can.setFillColor(HexColor("#0F172A"))
            
            text_str = f"Item: {product_name}"
            if qty and qty > 1:
                text_str += f"  (Qty: {qty})"
            if order_num:
                text_str = f"Order #{order_num}  |  {text_str}"

            # Left aligned text inside footer
            bg_can.drawString(8, 6, text_str[:100])
            bg_can.save()

            bg_buf.seek(0)
            bg_reader = PdfReader(bg_buf)
            new_page = bg_reader.pages[0]

            # Scale original page and shift it above the footer
            x_offset = width * (1.0 - scale) / 2.0
            page.add_transformation(Transformation().scale(scale, scale).translate(x_offset, footer_height))
            new_page.merge_page(page)

            writer.add_page(new_page)

        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()
    except Exception as e:
        import logging
        logging.getLogger("crm_api").error(f"Error stamping PDF: {e}")
        return pdf_bytes


@router.get("/{order_id}/download-label")
def download_order_label(
    order_id: int,
    download: bool = False,
    db: Session = Depends(get_db)
):
    """
    Downloads or views the order label PDF with the product name stamped at the bottom.
    Also names the file with tracking ID and product name.
    """
    order = db.query(Order).filter(Order.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    if not order.label_pdf_url:
        raise HTTPException(status_code=404, detail="No label PDF uploaded for this order")

    pdf_bytes = None
    pdf_url = order.label_pdf_url.strip()

    if pdf_url.startswith("http://") or pdf_url.startswith("https://"):
        import urllib.request
        req = urllib.request.Request(pdf_url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req) as resp:
            pdf_bytes = resp.read()
    else:
        BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        clean_rel = pdf_url.lstrip("/")
        if clean_rel.startswith("uploads/"):
            clean_rel = clean_rel[len("uploads/"):]
        file_path = os.path.join(BASE_DIR, "uploads", clean_rel)
        if not os.path.exists(file_path):
            raise HTTPException(status_code=404, detail="Label PDF file not found on disk")
        with open(file_path, "rb") as f:
            pdf_bytes = f.read()

    # Stamp the product name at the bottom of the PDF
    product_name = order.product_name or "Item"
    stamped_bytes = stamp_product_name_on_pdf(
        pdf_bytes=pdf_bytes,
        product_name=product_name,
        order_num=order.order_number or "",
        qty=order.qty or 1
    )

    clean_product = re.sub(r'[^a-zA-Z0-9_\-\. ]', '_', product_name)[:50].strip()
    tracking_part = order.label_tracking_id or order.shipment_id or order.order_number or f"Order-{order.id}"
    download_filename = f"{tracking_part} - {clean_product}.pdf"

    disposition = "attachment" if download else "inline"
    return Response(
        content=stamped_bytes,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'{disposition}; filename="{download_filename}"'
        }
    )



def parse_date_flexible(val: Any) -> Optional[date]:
    if val is None:
        return None
    if isinstance(val, date) and not isinstance(val, datetime):
        return val
    if isinstance(val, datetime):
        return val.date()
    cleaned = str(val).strip()
    if not cleaned or cleaned.lower() in ["none", "null", "—", "-", "n/a"]:
        return None
    if cleaned.endswith(" 00:00:00"):
        cleaned = cleaned[:-9].strip()
    for fmt in [
        "%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%m-%d-%Y",
        "%d/%m/%Y", "%d-%m-%Y", "%b %d %Y", "%b %d, %Y",
        "%B %d %Y", "%B %d, %Y", "%d %b %Y", "%d %B %Y",
        "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%m/%d/%Y %H:%M:%S",
        "%d.%m.%Y", "%Y.%m.%d"
    ]:
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            pass
    return None

def normalize_key(k: Any) -> str:
    return re.sub(r'[^a-z0-9]', '', str(k).lower().strip())

def parse_file_rows(filename: str, content_bytes: bytes) -> List[List[str]]:
    fname = filename.lower()
    if fname.endswith(".csv"):
        text = ""
        for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
            try:
                text = content_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if not text:
            raise HTTPException(status_code=400, detail="Could not decode CSV file. Please upload a valid UTF-8 CSV.")
        reader = csv.reader(io.StringIO(text))
        rows = []
        for r in reader:
            rows.append([str(c).strip() for c in r])
        return rows
    elif fname.endswith(".xlsx") or fname.endswith(".xls"):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
            sheet = wb.active
            rows = []
            for row in sheet.iter_rows(values_only=True):
                row_vals = []
                for val in row:
                    if val is None:
                        row_vals.append("")
                    elif isinstance(val, (datetime, date)):
                        row_vals.append(val.strftime("%Y-%m-%d"))
                    elif isinstance(val, float) and val.is_integer():
                        row_vals.append(str(int(val)))
                    else:
                        row_vals.append(str(val).strip())
                rows.append(row_vals)
            return rows
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Could not read Excel file: {str(e)}")
    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file format. Please upload a valid .csv, .xlsx, or .xls file."
        )

KEY_MAPPINGS = {
    "order_number": ["ordernumber", "orderid", "orderno", "order", "ordernum", "amazonorderid", "id", "orderidentifier"],
    "order_process_date": ["orderprocessdate", "processdate", "orderdate", "date", "purchasedate", "orderprocesseddate"],
    "shipping_date": ["shippingdate", "shipdate", "dateshipped", "dispatcheddate", "shipmentdate"],
    "last_delivery_date": ["lastdeliverydate", "deliverydate", "deliverby", "latestdeliverydate", "expecteddeliverydate", "maxdeliverydate"],
    "arriving_date": ["arrivingdate", "arrivaldate", "estimatedarrival", "expectedarrivaldate"],
    "company": ["company", "companyname", "sourcingagent", "client", "person"],
    "shipment_id": ["shipmentid", "shipmentnumber", "shipmentno", "awb", "trackingid", "shipmentcode"],
    "seller_account": ["selleraccount", "seller", "store", "account", "sellername", "merchant", "storeaccount"],
    "product_name": ["productname", "product", "title", "itemname", "item", "description", "producttitle"],
    "product_url": ["producturl", "url", "link", "itemurl", "asinurl", "productlink"],
    "product_image": ["productimage", "image", "imageurl", "img", "photourl"],
    "qty": ["qty", "quantity", "quantitypurchased", "units", "count", "itemqty"],
    "price_usd": ["priceusd", "price", "productprice", "itemprice", "unitprice", "amount", "totalprice", "itemsubtotal", "saleamount"],
    "order_status": ["orderstatus", "carrierstatus", "stockstatus", "status"],
    "consignee_name": ["consigneename", "consignee", "buyername", "buyer", "customername", "customer", "recipientname", "recipient", "name"],
    "shipment_address_1": ["shipmentaddress1", "address1", "shipaddress1", "streetaddress", "address", "street", "addressline1"],
    "shipment_address_2": ["shipmentaddress2", "address2", "shipaddress2", "apartment", "suite", "unit", "addressline2"],
    "city": ["city", "shipcity", "destinationcity"],
    "state": ["state", "shipstate", "destinationstate", "province", "region"],
    "zip_code": ["zipcode", "zip", "shipzip", "postalcode", "pincode", "postcode"],
    "mobile_number": ["mobilenumber", "mobile", "phone", "phonenumber", "shipphone", "contactnumber", "contactphone"],
    "country": ["country", "shipcountry", "destinationcountry"],
    "purchase_cost_inr": ["purchasecostinr", "purchasecost", "cost", "cogs", "itemcost"],
    "shipment_cost": ["shipmentcost", "shippingcost", "shippingfee", "deliverycost"],
}

SAMPLE_HEADERS = [
    "Order ID",
    "Order Process Date",
    "Shipping Date",
    "Last Delivery Date",
    "Arriving Date",
    "Company / Person",
    "Shipment ID",
    "Seller Account",
    "Product Name",
    "Product URL",
    "Product Image URL",
    "Qty",
    "Price ($)",
    "Order Status",
    "Consignee Name",
    "Address Line 1",
    "Address Line 2",
    "City",
    "State",
    "Zip Code",
    "Contact Number",
    "Country",
    "Purchase Cost (INR)",
    "Shipment Cost"
]

SAMPLE_DATA_ROWS = [
    [
        "114-1029384-5928173",
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        "ADBH",
        "INBTL001",
        "Store US #1",
        "Sea Buckthorn Berry Juice 500ml",
        "https://example.com/product/sea-buckthorn",
        "",
        "1",
        "34.99",
        "ADBH",
        "John Doe",
        "123 Maple Street",
        "Apt 4B",
        "New York",
        "NY",
        "10001",
        "+1 555-0199",
        "USA",
        "1250.00",
        "15.00"
    ],
    [
        "114-5544332-2211009",
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        "Vetai",
        "INBTL002",
        "Store US #2",
        "Organic Vitamin C Serum 30ml | Hydrating Hyaluronic Cream 50g",
        "https://example.com/product/serum | https://example.com/product/cream",
        "",
        "2 | 1",
        "19.99 | 24.50",
        "ADBH",
        "Alice Smith",
        "456 Oak Avenue",
        "Suite 100",
        "Chicago",
        "IL",
        "60601",
        "+1 555-0188",
        "USA",
        "1450.00",
        "12.00"
    ],
    [
        "114-8849201-9482019",
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        "Globle",
        "INBTL003",
        "Store US #3",
        "Natural Tea Tree Face Wash 150ml",
        "https://example.com/product/face-wash",
        "",
        "2",
        "18.00",
        "ADBH",
        "Sarah Connor",
        "742 Evergreen Terrace",
        "",
        "Los Angeles",
        "CA",
        "90001",
        "+1 555-0144",
        "USA",
        "700.00",
        "10.00"
    ],
    [
        "114-8849201-9482019",
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        date.today().strftime("%Y-%m-%d"),
        "Globle",
        "INBTL003",
        "Store US #3",
        "Purifying Charcoal Clay Mask 100g",
        "https://example.com/product/clay-mask",
        "",
        "1",
        "22.50",
        "ADBH",
        "Sarah Connor",
        "742 Evergreen Terrace",
        "",
        "Los Angeles",
        "CA",
        "90001",
        "+1 555-0144",
        "USA",
        "800.00",
        "0.00"
    ]
]

@router.get("/sample-template")
def download_sample_template(
    format: str = Query("csv", pattern="^(csv|xlsx)$"),
    current_user: User = Depends(check_permission("orders:read"))
):
    if format == "csv":
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(SAMPLE_HEADERS)
        for row in SAMPLE_DATA_ROWS:
            writer.writerow(row)
        csv_bytes = out.getvalue().encode("utf-8-sig")
        return Response(
            content=csv_bytes,
            media_type="text/csv",
            headers={
                "Content-Disposition": "attachment; filename=orders_sample_template.csv",
                "Access-Control-Expose-Headers": "Content-Disposition"
            }
        )
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Orders Import Template"

        # Styles
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=11)
        note_font = Font(name="Calibri", size=10, italic=True, color="64748B")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        ws.append(SAMPLE_HEADERS)

        for col_idx in range(1, len(SAMPLE_HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for row_idx, row_data in enumerate(SAMPLE_DATA_ROWS, start=2):
            ws.append(row_data)
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if col_idx in [12]: # Qty
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx in [13, 23, 24]: # Prices
                    cell.alignment = Alignment(horizontal="right")

        # Instructions sheet
        ws_info = wb.create_sheet(title="Import Guide & Rules")
        ws_info.column_dimensions['A'].width = 30
        ws_info.column_dimensions['B'].width = 75

        info_headers = ["Field / Rule", "Description & Required Format"]
        ws_info.append(info_headers)
        ws_info.cell(row=1, column=1).font = header_font
        ws_info.cell(row=1, column=1).fill = header_fill
        ws_info.cell(row=1, column=2).font = header_font
        ws_info.cell(row=1, column=2).fill = header_fill

        rules = [
            ("Order ID (Required)", "Unique Order ID (e.g. 114-1029384-5928173). Required on every row."),
            ("Multi-Product Orders", "To create a multi-item order, put multiple rows with the SAME Order ID. They will be combined into a single order with individual product items."),
            ("Product Name (Required)", "Name/title of the product. Cannot be empty."),
            ("Qty (Required)", "Integer quantity >= 1 (e.g. 1, 2, 5)."),
            ("Price ($) (Required)", "Sale price in USD (e.g. 29.99, 15.00). Must be non-negative."),
            ("Date Fields", "Order Process Date, Shipping Date, Last Delivery Date, Arriving Date. Expected format: YYYY-MM-DD or DD/MM/YYYY."),
            ("Company / Person", "Must match your allowed company (e.g. ADBH, Vetai, Globle)."),
            ("Atomic Validation", "All rows are validated strictly. If any row has a missing required field or error, the whole file is rejected with clear error details."),
        ]

        for r_idx, (k, v) in enumerate(rules, start=2):
            ws_info.append([k, v])
            ws_info.cell(row=r_idx, column=1).font = Font(bold=True)
            ws_info.cell(row=r_idx, column=1).border = thin_border
            ws_info.cell(row=r_idx, column=2).border = thin_border

        # Adjust column widths in main sheet
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        ws.row_dimensions[1].height = 28

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": "attachment; filename=orders_sample_template.xlsx",
                "Access-Control-Expose-Headers": "Content-Disposition"
            }
        )

@router.post("/upload-bulk-file")
@router.post("/upload-csv")
async def upload_bulk_orders_file(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("orders:write"))
):
    content_bytes = await file.read()
    if not content_bytes:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    rows = parse_file_rows(file.filename, content_bytes)
    
    # Filter out empty rows
    data_rows = []
    for r in rows:
        if any(str(c).strip() for c in r):
            data_rows.append(r)

    if not data_rows or len(data_rows) < 2:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file contains no data rows. Please ensure the first row has column headers and subsequent rows have order data."
        )

    header_row = data_rows[0]
    header_indices: dict[str, int] = {}

    for idx, col in enumerate(header_row):
        norm = normalize_key(col)
        for std_key, aliases in KEY_MAPPINGS.items():
            if norm in aliases or norm == std_key:
                if std_key not in header_indices:
                    header_indices[std_key] = idx
                break

    # Verify essential columns exist in headers
    missing_required_headers = []
    if "order_number" not in header_indices:
        missing_required_headers.append("Order ID (or Order Number)")
    if "product_name" not in header_indices:
        missing_required_headers.append("Product Name")

    if missing_required_headers:
        raise HTTPException(
            status_code=400,
            detail=f"Required column header(s) missing: {', '.join(missing_required_headers)}. Please download the sample template for the correct header format."
        )

    allowed_comps = get_user_allowed_companies(current_user)
    acc_id = current_user.account_id if not current_user.is_admin else None
    acc_name = current_user.account_name if not current_user.is_admin else None

    # Strict Validation Collection
    validation_errors: List[str] = []
    
    # Group rows by order_number
    # Format: { order_number: [ { "row_idx": int, "data": dict }, ... ] }
    grouped_orders: dict[str, List[dict]] = {}
    
    # Track order numbers seen in file to catch duplicate IDs in DB
    existing_db_orders = set()
    order_numbers_in_file = set()

    for row_idx, row in enumerate(data_rows[1:], start=2):
        def get_val(key: str, default="") -> str:
            idx = header_indices.get(key)
            if idx is not None and idx < len(row):
                val = str(row[idx]).strip()
                return val if val else default
            return default

        raw_order_num = get_val("order_number")
        if not raw_order_num:
            validation_errors.append(f"Row {row_idx}: 'Order ID' is missing or empty.")
            continue

        raw_product_name = get_val("product_name")
        if not raw_product_name:
            validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): 'Product Name' is required.")
            continue

        product_parts = [p.strip() for p in re.split(r'\||\r?\n', raw_product_name) if p.strip()]

        # Validate Optional Costs
        raw_purchase_cost = get_val("purchase_cost_inr", "0")
        try:
            purchase_cost_val = float(str(raw_purchase_cost).replace("₹", "").replace(",", "").strip())
            if purchase_cost_val < 0:
                validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): 'Purchase Cost' cannot be negative (got '{raw_purchase_cost}').")
        except (ValueError, TypeError):
            validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): 'Purchase Cost' must be a valid number (got '{raw_purchase_cost}').")
            purchase_cost_val = 0.0

        raw_shipment_cost = get_val("shipment_cost", "0")
        try:
            shipment_cost_val = float(str(raw_shipment_cost).replace("$", "").replace(",", "").strip())
            if shipment_cost_val < 0:
                validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): 'Shipment Cost' cannot be negative (got '{raw_shipment_cost}').")
        except (ValueError, TypeError):
            validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): 'Shipment Cost' must be a valid number (got '{raw_shipment_cost}').")
            shipment_cost_val = 0.0

        # Validate Dates
        raw_process_date = get_val("order_process_date")
        process_date = None
        if raw_process_date:
            process_date = parse_date_flexible(raw_process_date)
            if not process_date:
                validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): Invalid 'Order Process Date' ('{raw_process_date}'). Expected YYYY-MM-DD or DD/MM/YYYY.")
        else:
            process_date = date.today()

        raw_shipping_date = get_val("shipping_date")
        shipping_date = None
        if raw_shipping_date:
            shipping_date = parse_date_flexible(raw_shipping_date)
            if not shipping_date:
                validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): Invalid 'Shipping Date' ('{raw_shipping_date}'). Expected YYYY-MM-DD or DD/MM/YYYY.")

        raw_last_delivery = get_val("last_delivery_date")
        last_delivery_date = None
        if raw_last_delivery:
            last_delivery_date = parse_date_flexible(raw_last_delivery)
            if not last_delivery_date:
                validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): Invalid 'Last Delivery Date' ('{raw_last_delivery}'). Expected YYYY-MM-DD or DD/MM/YYYY.")

        raw_arriving_date = get_val("arriving_date")
        arriving_date = None
        if raw_arriving_date:
            arriving_date = parse_date_flexible(raw_arriving_date)
            if not arriving_date and not (len(raw_arriving_date) <= 30 and ("-" in raw_arriving_date or "/" in raw_arriving_date)):
                validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): Invalid 'Arriving Date' ('{raw_arriving_date}'). Expected YYYY-MM-DD or DD/MM/YYYY.")

        # Validate Company Authorization
        company = get_val("company", "ADBH")
        if allowed_comps:
            comp_clean = company.strip().lower()
            if comp_clean and not any(c.lower() in comp_clean for c in allowed_comps):
                validation_errors.append(
                    f"Row {row_idx} (Order ID '{raw_order_num}'): Not authorized to create orders for company '{company}'. Allowed: {', '.join(allowed_comps)}."
                )

        order_numbers_in_file.add(raw_order_num)

        # Check if single row has multiple products (e.g. "Product A | Product B" or separated by newlines)
        if len(product_parts) > 1:
            raw_qty_str = get_val("qty", "1")
            qty_parts = [q.strip() for q in re.split(r'\||\r?\n', raw_qty_str) if q.strip()]
            raw_price_str = get_val("price_usd", "0")
            price_parts = [p.strip() for p in re.split(r'\||\r?\n', raw_price_str) if p.strip()]
            url_parts = [u.strip() for u in re.split(r'\||\r?\n', get_val("product_url")) if u.strip()]
            img_parts = [img.strip() for img in re.split(r'\||\r?\n', get_val("product_image")) if img.strip()]

            for p_idx, p_name in enumerate(product_parts):
                # Qty for sub-product
                if p_idx < len(qty_parts):
                    cur_qty_str = qty_parts[p_idx]
                elif len(qty_parts) == 1:
                    cur_qty_str = qty_parts[0]
                else:
                    cur_qty_str = "1"

                try:
                    p_qty_val = int(float(cur_qty_str))
                    if p_qty_val < 1:
                        validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}', Product #{p_idx+1} '{p_name}'): 'Qty' must be at least 1 (got '{cur_qty_str}').")
                except (ValueError, TypeError):
                    validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}', Product #{p_idx+1} '{p_name}'): 'Qty' must be a valid integer (got '{cur_qty_str}').")
                    p_qty_val = 1

                # Price for sub-product
                if p_idx < len(price_parts):
                    cur_price_str = price_parts[p_idx]
                elif len(price_parts) == 1 and p_idx == 0:
                    cur_price_str = price_parts[0]
                else:
                    cur_price_str = "0.0"

                try:
                    p_price_val = float(str(cur_price_str).replace("$", "").replace(",", "").strip())
                    if p_price_val < 0:
                        validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}', Product #{p_idx+1} '{p_name}'): 'Price ($)' cannot be negative (got '{cur_price_str}').")
                except (ValueError, TypeError):
                    validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}', Product #{p_idx+1} '{p_name}'): 'Price ($)' must be a valid numeric amount (got '{cur_price_str}').")
                    p_price_val = 0.0

                p_url = url_parts[p_idx] if p_idx < len(url_parts) else (url_parts[0] if len(url_parts) == 1 else "")
                p_img = img_parts[p_idx] if p_idx < len(img_parts) else (img_parts[0] if len(img_parts) == 1 else "")

                sub_item = {
                    "row_idx": row_idx,
                    "order_number": raw_order_num,
                    "order_process_date": process_date,
                    "shipping_date": shipping_date,
                    "last_delivery_date": last_delivery_date,
                    "arriving_date": raw_arriving_date,
                    "company": company,
                    "shipment_id": get_val("shipment_id"),
                    "seller_account": get_val("seller_account", acc_name or ""),
                    "product_name": p_name,
                    "product_url": p_url,
                    "product_image": p_img,
                    "qty": p_qty_val,
                    "price_usd": p_price_val,
                    "order_status": get_val("order_status", "ADBH"),
                    "consignee_name": get_val("consignee_name", "Consignee"),
                    "shipment_address_1": get_val("shipment_address_1"),
                    "shipment_address_2": get_val("shipment_address_2"),
                    "city": get_val("city"),
                    "state": get_val("state"),
                    "zip_code": get_val("zip_code"),
                    "mobile_number": get_val("mobile_number"),
                    "country": get_val("country", "USA"),
                    "purchase_cost_inr": purchase_cost_val if p_idx == 0 else 0.0,
                    "shipment_cost": shipment_cost_val if p_idx == 0 else 0.0
                }

                if raw_order_num not in grouped_orders:
                    grouped_orders[raw_order_num] = []
                grouped_orders[raw_order_num].append(sub_item)

        else:
            # Single product in row
            raw_qty = get_val("qty", "1")
            try:
                qty_val = int(float(raw_qty))
                if qty_val < 1:
                    validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): 'Qty' must be at least 1 (got '{raw_qty}').")
            except (ValueError, TypeError):
                validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): 'Qty' must be a valid integer (got '{raw_qty}').")
                qty_val = 1

            raw_price = get_val("price_usd", "0")
            try:
                price_val = float(str(raw_price).replace("$", "").replace(",", "").strip())
                if price_val < 0:
                    validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): 'Price ($)' cannot be negative (got '{raw_price}').")
            except (ValueError, TypeError):
                validation_errors.append(f"Row {row_idx} (Order ID '{raw_order_num}'): 'Price ($)' must be a valid numeric amount (got '{raw_price}').")
                price_val = 0.0

            row_item = {
                "row_idx": row_idx,
                "order_number": raw_order_num,
                "order_process_date": process_date,
                "shipping_date": shipping_date,
                "last_delivery_date": last_delivery_date,
                "arriving_date": raw_arriving_date,
                "company": company,
                "shipment_id": get_val("shipment_id"),
                "seller_account": get_val("seller_account", acc_name or ""),
                "product_name": product_parts[0],
                "product_url": get_val("product_url"),
                "product_image": get_val("product_image"),
                "qty": qty_val,
                "price_usd": price_val,
                "order_status": get_val("order_status", "ADBH"),
                "consignee_name": get_val("consignee_name", "Consignee"),
                "shipment_address_1": get_val("shipment_address_1"),
                "shipment_address_2": get_val("shipment_address_2"),
                "city": get_val("city"),
                "state": get_val("state"),
                "zip_code": get_val("zip_code"),
                "mobile_number": get_val("mobile_number"),
                "country": get_val("country", "USA"),
                "purchase_cost_inr": purchase_cost_val,
                "shipment_cost": shipment_cost_val
            }

            if raw_order_num not in grouped_orders:
                grouped_orders[raw_order_num] = []
            grouped_orders[raw_order_num].append(row_item)

    # Check for existing duplicate Order IDs in database
    if order_numbers_in_file:
        existing_orders = db.query(Order.order_number).filter(Order.order_number.in_(list(order_numbers_in_file))).all()
        for (exist_num,) in existing_orders:
            first_row_idx = grouped_orders[exist_num][0]["row_idx"] if exist_num in grouped_orders else "Unknown"
            validation_errors.append(f"Row {first_row_idx}: Order ID '{exist_num}' already exists in the system. Duplicate order IDs are not allowed.")

    # ZERO TOLERANCE: If ANY error occurred, abort entire transaction and return errors
    if validation_errors:
        raise HTTPException(
            status_code=422,
            detail={
                "message": f"Validation failed with {len(validation_errors)} error(s). No orders were processed.",
                "error_count": len(validation_errors),
                "errors": validation_errors
            }
        )

    # All validations passed! Proceed with atomic creation
    created_orders = []
    total_items_count = 0
    months_to_sync = set()

    try:
        for order_num, items in grouped_orders.items():
            first_item = items[0]
            total_qty = sum(it["qty"] for it in items)
            total_price = sum(it["price_usd"] for it in items)
            total_purchase_cost = sum(it["purchase_cost_inr"] for it in items)
            total_shipment_cost = max(it["shipment_cost"] for it in items)
            total_items_count += len(items)

            # Combined product name
            combined_product_name = " | ".join(it["product_name"] for it in items)

            # Build product_items array
            product_items_list = []
            primary_inv_id = None
            primary_image = first_item["product_image"]
            primary_url = first_item["product_url"]

            for it in items:
                inv = get_or_create_inventory_item(
                    db=db,
                    product_name=it["product_name"],
                    price_usd=it["price_usd"],
                    seller_account=it["seller_account"] or it["company"],
                    image_url=it["product_image"],
                    product_url=it["product_url"]
                )
                if inv and primary_inv_id is None:
                    primary_inv_id = inv.id
                
                it_image = it["product_image"] or (inv.image_url if inv else None)
                it_url = it["product_url"] or (inv.product_url if inv else None)
                
                if not primary_image and it_image:
                    primary_image = it_image
                if not primary_url and it_url:
                    primary_url = it_url

                product_items_list.append({
                    "product_id": inv.id if inv else None,
                    "product_name": it["product_name"],
                    "product_url": it_url,
                    "product_image": it_image,
                    "qty": it["qty"],
                    "price_usd": it["price_usd"]
                })

            product_items_json = json.dumps(product_items_list) if len(product_items_list) > 1 else None

            shipment_num = first_item["shipment_id"] or get_next_shipment_id(db, offset=len(created_orders))

            new_order = Order(
                order_number=order_num,
                order_process_date=first_item["order_process_date"] or date.today(),
                shipping_date=first_item["shipping_date"],
                last_delivery_date=first_item["last_delivery_date"],
                arriving_date=first_item["arriving_date"],
                company=first_item["company"],
                shipment_id=shipment_num,
                seller_account=first_item["seller_account"],
                product_id=primary_inv_id,
                product_name=combined_product_name,
                product_url=primary_url,
                product_image=primary_image,
                product_items=product_items_json,
                qty=total_qty,
                product_price=total_price,
                order_status=first_item["order_status"] or "ADBH",
                purchase_cost_inr=total_purchase_cost,
                shipment_cost=total_shipment_cost,
                consignee_name=first_item["consignee_name"] or "Consignee",
                shipment_address_1=first_item["shipment_address_1"] or "",
                shipment_address_2=first_item["shipment_address_2"] or "",
                city=first_item["city"] or "",
                state=first_item["state"] or "",
                zip_code=first_item["zip_code"] or "",
                mobile_number=first_item["mobile_number"] or "",
                country=first_item["country"] or "USA",
                account_id=acc_id,
                account_name=acc_name,
                status="Pending"
            )
            db.add(new_order)
            created_orders.append(new_order)

            if first_item["order_process_date"]:
                months_to_sync.add(first_item["order_process_date"].strftime("%Y-%m"))

        db.commit()

        for m in months_to_sync:
            sync_admin_cost_share_for_month(m, db)

        return {
            "success": True,
            "message": f"Successfully imported {len(created_orders)} order(s) with {total_items_count} product item(s).",
            "imported_orders_count": len(created_orders),
            "imported_items_count": total_items_count
        }

    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail=f"An error occurred while saving imported orders: {str(exc)}"
        )

@router.post("/bulk", response_model=List[OrderResponse])
def create_bulk_orders(
    orders_in: List[OrderCreate],
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("orders:write"))
):
    allowed_comps = get_user_allowed_companies(current_user)
    if allowed_comps:
        for o_chk in orders_in:
            comp_val = (o_chk.company or "").strip().lower()
            if comp_val and not any(c.lower() in comp_val for c in allowed_comps):
                raise HTTPException(
                    status_code=403,
                    detail=f"Not authorized to create orders for company '{o_chk.company}'. Allowed: {', '.join(allowed_comps)}"
                )

    created_orders = []
    acc_id = current_user.account_id if not current_user.is_admin else None
    acc_name = current_user.account_name if not current_user.is_admin else None

    # Determine a base shipment id if none provided across batch
    common_shipment_num = None
    first_shipment = orders_in[0].shipment_id if orders_in else None

    for idx, order_in in enumerate(orders_in):
        inventory_item = None
        if order_in.product_id:
            inventory_item = db.query(Inventory).filter(Inventory.id == order_in.product_id).first()
        if not inventory_item and order_in.product_name:
            inventory_item = get_or_create_inventory_item(
                db=db,
                product_name=order_in.product_name,
                price_usd=order_in.price_usd or 0.0,
                seller_account=order_in.seller_account or order_in.company,
                image_url=order_in.product_image,
                product_url=order_in.product_url
            )
        elif inventory_item and order_in.product_url and order_in.product_url.strip():
            if not inventory_item.product_url:
                inventory_item.product_url = order_in.product_url.strip()
                db.commit()

        effective_product_url = (order_in.product_url or "").strip() or (inventory_item.product_url if inventory_item else None)
        effective_product_image = order_in.product_image or (inventory_item.image_url if inventory_item else None)
        if not effective_product_image and effective_product_url:
            extracted = extract_product_info_from_url(effective_product_url)
            if extracted.get("image_url"):
                effective_product_image = extracted["image_url"]
                if inventory_item and not inventory_item.image_url:
                    inventory_item.image_url = effective_product_image
                    db.commit()

        order_num = order_in.order_number or f"114-{random.randint(1000000, 9999999)}-{random.randint(1000000, 9999999)}"
        shipment_num = order_in.shipment_id or first_shipment or get_next_shipment_id(db, offset=idx)

        item_acc_id = acc_id or order_in.account_id
        item_acc_name = acc_name or order_in.account_name
        if not item_acc_id and order_in.seller_account:
            acc = db.query(Account).filter(Account.account_name.ilike(order_in.seller_account.strip())).first()
            if acc:
                item_acc_id = acc.id
                item_acc_name = acc.account_name

        order = Order(
            order_number=order_num,
            order_process_date=order_in.order_process_date or date.today(),
            last_delivery_date=order_in.last_delivery_date,
            shipping_date=order_in.shipping_date,
            company=order_in.company,
            shipment_id=shipment_num,
            seller_account=order_in.seller_account or "",
            product_id=inventory_item.id if inventory_item else None,
            product_name=order_in.product_name or f"Item #{idx + 1}",
            product_url=effective_product_url,
            product_image=effective_product_image,
            qty=order_in.qty or 1,
            product_price=getattr(order_in, 'product_price', None) or order_in.price_usd or 0.0,
            order_status=order_in.order_status or "ADBH",
            purchase_cost_inr=order_in.purchase_cost_inr or 0.0,
            arriving_date=order_in.arriving_date,
            consignee_name=order_in.consignee_name or "Consignee",
            shipment_address_1=order_in.shipment_address_1 or "",
            shipment_address_2=order_in.shipment_address_2 or "",
            city=order_in.city or "",
            state=order_in.state or "",
            zip_code=order_in.zip_code or "",
            mobile_number=order_in.mobile_number or "",
            country=order_in.country or "USA",
            account_id=item_acc_id,
            account_name=item_acc_name,
            status=order_in.status or "Pending",
            delivery_service=order_in.delivery_service,
            shipment_cost=order_in.shipment_cost or 0.0
        )
        db.add(order)
        created_orders.append(order)

    db.commit()
    
    months = set()
    for o in created_orders:
        db.refresh(o)
        d = o.order_date or o.order_process_date or date.today()
        if d:
            months.add(d.strftime("%Y-%m"))
            
    for m in months:
        sync_admin_cost_share_for_month(m, db)
        
    for o in created_orders:
        db.refresh(o)
        
    populate_order_costs(created_orders, db)
    return created_orders

def check_order_access(order: Order, current_user: User):
    if not current_user.is_admin:
        allowed_comps = get_user_allowed_companies(current_user)
        if allowed_comps:
            ord_comp = (order.company or "").strip().lower()
            if not any(c.lower() in ord_comp for c in allowed_comps):
                raise HTTPException(status_code=403, detail="Not authorized to access this company order")
        elif current_user.account_name:
            comp_name = current_user.account_name.strip().lower()
            ord_comp = (order.company or "").strip().lower()
            ord_acc = (order.account_name or "").strip().lower()
            if comp_name not in ord_comp and comp_name not in ord_acc:
                raise HTTPException(status_code=403, detail="Not authorized to access this company order")

@router.get("/extract-url-image")
def extract_url_image(
    url: str = Query(..., description="Product or Image URL to extract preview image from"),
    current_user: User = Depends(get_current_user)
):
    return extract_product_info_from_url(url)

@router.get("/{order_id}", response_model=OrderResponse)
def get_order(
    order_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("orders:read"))
):
    order = db.query(Order).filter(Order.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    check_order_access(order, current_user)
    populate_order_costs(order, db)
    return order

@router.put("/{order_id}", response_model=OrderResponse)
def update_order(
    order_id: int,
    order_in: OrderUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("orders:write"))
):
    order = db.query(Order).filter(Order.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    check_order_access(order, current_user)

    old_date = order.order_date or order.order_process_date or date.today()
    old_month = old_date.strftime("%Y-%m") if old_date else None

    old_status = order.status

    for field, value in order_in.dict(exclude_unset=True).items():
        if field == "price_usd" and value is not None:
            order.product_price = value
        elif hasattr(order, field):
            setattr(order, field, value)

    if order_in.seller_account:
        acc = db.query(Account).filter(Account.account_name.ilike(order_in.seller_account.strip())).first()
        if acc:
            order.account_id = acc.id
            order.account_name = acc.account_name

    # Ensure product_id is linked if missing
    if not order.product_id and order.product_name:
        inv = db.query(Inventory).filter(
            func.lower(func.trim(Inventory.product_name)) == order.product_name.strip().lower()
        ).first()
        if inv:
            order.product_id = inv.id

    # Sync product_url with inventory
    if order_in.product_url and order_in.product_url.strip():
        clean_url = order_in.product_url.strip()
        order.product_url = clean_url
        if order.product_id:
            inv = db.query(Inventory).filter(Inventory.id == order.product_id).first()
            if inv and (not inv.product_url or inv.product_url != clean_url):
                inv.product_url = clean_url
    elif (not order.product_url or not order.product_url.strip()) and order.product_id:
        inv = db.query(Inventory).filter(Inventory.id == order.product_id).first()
        if inv and inv.product_url:
            order.product_url = inv.product_url

    if order_in.status is not None and order_in.status == "Ready to Ship" and old_status != "Ready to Ship":
        inv = None
        if order.product_id:
            inv = db.query(Inventory).filter(Inventory.id == order.product_id).first()
        else:
            inv = db.query(Inventory).filter(Inventory.product_name.ilike(order.product_name)).first()
        if inv:
            inv.stock_quantity = max(0, inv.stock_quantity - order.qty)

    db.commit()
    db.refresh(order)

    new_date = order.order_date or order.order_process_date or date.today()
    new_month = new_date.strftime("%Y-%m") if new_date else None

    if old_month:
        sync_admin_cost_share_for_month(old_month, db)
    if new_month and new_month != old_month:
        sync_admin_cost_share_for_month(new_month, db)

    db.refresh(order)
    populate_order_costs(order, db)
    return order

@router.delete("/{order_id}")
def delete_order(
    order_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("orders:write"))
):
    order = db.query(Order).filter(Order.id == order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    check_order_access(order, current_user)

    d = order.order_date or order.order_process_date or date.today()
    month = d.strftime("%Y-%m") if d else None

    from app.core.s3 import delete_file_from_s3
    if order.label_pdf_url:
        delete_file_from_s3(order.label_pdf_url)

    db.delete(order)
    db.commit()

    if month:
        sync_admin_cost_share_for_month(month, db)

    return {"message": "Order deleted successfully"}
